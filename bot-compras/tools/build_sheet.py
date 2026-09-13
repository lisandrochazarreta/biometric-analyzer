#!/usr/bin/env python3
"""Genera el .xlsx 'Compras Casa' que se sube a Drive convertido a Google Sheet."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

CATS = ['almacén', 'lácteos', 'limpieza', 'higiene', 'bebidas', 'frescos', 'otros']
ESTADOS = ['pendiente', 'comprado', 'descartado', 'archivado']

HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(color='FFFFFF', bold=True, size=11)

def sheet(wb, title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 22
    return ws

wb = Workbook()

# ---------------- Lista ----------------
LISTA_H = ['id','ciclo','item','variantes','cantidad','unidad','categoria','pedido_por',
           'nota','fecha_alta','estado','fecha_compra','veces_arrastrado','origen','msg_origen']
ws = sheet(wb, 'Lista', LISTA_H, [22,10,26,30,9,8,12,14,20,22,12,22,9,10,18], first=True)
ws.auto_filter.ref = f'A1:{get_column_letter(len(LISTA_H))}1'

dv_cat = DataValidation(type='list', formula1='"%s"' % ','.join(CATS), allow_blank=True)
dv_cat.error = 'Categoría no válida'
ws.add_data_validation(dv_cat); dv_cat.add('G2:G2000')

dv_est = DataValidation(type='list', formula1='"%s"' % ','.join(ESTADOS), allow_blank=True)
dv_est.error = 'Estado no válido'
ws.add_data_validation(dv_est); dv_est.add('K2:K2000')

rng = 'A2:O2000'
# comprado -> gris tachado
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['$K2="comprado"'],
    font=Font(color='9AA0A6', strike=True), stopIfTrue=False))
# descartado -> gris claro
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['$K2="descartado"'],
    font=Font(color='C0C4C8', strike=True), stopIfTrue=False))
# arrastrado y pendiente -> amarillo
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['AND($K2="pendiente",$M2>0)'],
    fill=PatternFill('solid', fgColor='FFF3CD'), stopIfTrue=False))
# a punto de descartarse
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['AND($K2="pendiente",$M2>=2)'],
    fill=PatternFill('solid', fgColor='F8D7DA'), stopIfTrue=False))

# ---------------- Historial ----------------
sheet(wb, 'Historial', LISTA_H + ['cerrado_el','cerrado_por'],
      [22,10,26,30,9,8,12,14,20,22,12,22,9,10,18,22,14])

# ---------------- Catalogo ----------------
cat = sheet(wb, 'Catalogo', ['item_canonico','categoria','alias','frecuencia','ultima_vez'],
            [26,12,60,11,14])
dv_cat2 = DataValidation(type='list', formula1='"%s"' % ','.join(CATS), allow_blank=True)
cat.add_data_validation(dv_cat2); dv_cat2.add('B2:B1000')

SEED = [
 ('Leche','lácteos','leche,lechee,leche entera,leche descremada,litro de leche,la leche'),
 ('Yogur','lácteos','yogur,yoghurt,yogures,yogurt'),
 ('Queso','lácteos','queso,queso cremoso,cremoso,queso de maquina,port salut'),
 ('Queso rallado','lácteos','rallado,queso rallado,parmesano'),
 ('Manteca','lácteos','manteca'),
 ('Crema','lácteos','crema,crema de leche'),
 ('Dulce de leche','lácteos','dulce de leche,ddl,dulce'),
 ('Café','almacén','cafe,café,cafe molido,cafe instantaneo,nescafe'),
 ('Yerba','almacén','yerba,yerba mate,la yerba'),
 ('Té','almacén','te,té,saquitos de te'),
 ('Azúcar','almacén','azucar,azúcar'),
 ('Harina','almacén','harina,harina 000,harina leudante'),
 ('Fideos','almacén','fideos,pasta,tallarines,mostachol,tirabuzon,spaghetti'),
 ('Arroz','almacén','arroz'),
 ('Aceite','almacén','aceite,aceite de girasol,aceite de oliva'),
 ('Vinagre','almacén','vinagre'),
 ('Sal','almacén','sal,sal fina,sal gruesa'),
 ('Puré de tomate','almacén','pure de tomate,puré de tomate,salsa de tomate,tomate triturado'),
 ('Atún','almacén','atun,atún,lata de atun'),
 ('Galletitas','almacén','galletitas,galletas,criollitas,vainillas'),
 ('Mermelada','almacén','mermelada,dulce de frutilla'),
 ('Polenta','almacén','polenta'),
 ('Lentejas','almacén','lentejas,legumbres,garbanzos'),
 ('Detergente','limpieza','detergente,deter,detergente para los platos,magistral'),
 ('Lavandina','limpieza','lavandina,ayudin,cloro'),
 ('Jabón en polvo','limpieza','jabon en polvo,jabón en polvo,skip,ala,jabon para la ropa'),
 ('Suavizante','limpieza','suavizante,vivere,comfort'),
 ('Esponjas','limpieza','esponjas,esponja,virulana'),
 ('Trapo de piso','limpieza','trapo,trapo de piso,trapos'),
 ('Limpiavidrios','limpieza','limpiavidrios,cif vidrios'),
 ('Bolsas de residuo','limpieza','bolsas,bolsas de residuo,bolsas de basura,bolsas de consorcio'),
 ('Papel de cocina','limpieza','papel de cocina,rollo de cocina,servilletas'),
 ('Papel higiénico','higiene','papel higienico,papel higiénico,papel,papel h,rollos,higienico'),
 ('Shampoo','higiene','shampoo,champu,champú,shampu'),
 ('Acondicionador','higiene','acondicionador,enjuague'),
 ('Jabón de tocador','higiene','jabon de tocador,jabon blanco,jaboncito,jabon de manos'),
 ('Pasta de dientes','higiene','pasta de dientes,dentifrico,colgate,pasta dental'),
 ('Desodorante','higiene','desodorante,deo,rexona'),
 ('Maquinitas de afeitar','higiene','maquinitas,gillette,maquinita de afeitar'),
 ('Agua','bebidas','agua,agua mineral,bidon de agua,agua sin gas'),
 ('Gaseosa','bebidas','gaseosa,coca,coca cola,sprite,seven up'),
 ('Cerveza','bebidas','cerveza,birra,porron,porrones'),
 ('Vino','bebidas','vino,vino tinto,malbec'),
 ('Fernet','bebidas','fernet,branca'),
 ('Jugo','bebidas','jugo,jugo de naranja,exprimido'),
 ('Soda','bebidas','soda,sifon'),
 ('Huevos','frescos','huevos,huevo,docena de huevos,maple de huevos'),
 ('Pan','frescos','pan,pan lactal,pan de mesa,flautita'),
 ('Pollo','frescos','pollo,pechuga,pata muslo'),
 ('Carne','frescos','carne,carne picada,milanesas,nalga,bife'),
 ('Verduras','frescos','verduras,verdura,lechuga,tomate,cebolla,papa,zanahoria'),
 ('Frutas','frescos','frutas,fruta,banana,manzana,naranja,limon'),
 ('Fiambre','frescos','fiambre,jamon,jamón,salame,queso y jamon'),
 ('Pilas','otros','pilas,pila,baterias,pilas AA'),
 ('Lamparitas','otros','lamparitas,lamparita,foco,bombita'),
 ('Comida para gato','otros','comida de gato,comida para gato,alimento para gato,whiskas'),
]
for item, c, alias in SEED:
    cat.append([item, c, alias, 0, ''])

# ---------------- Config ----------------
cfg = sheet(wb, 'Config', ['clave','valor','descripcion'], [34, 42, 62])
CONFIG = [
 ('ciclo_activo','2026-10','Mes en el que se compra. El cron del dia 1 lo avanza solo.'),
 ('grupo_chat_id','<PEGAR>','chat_id del grupo de Telegram (negativo, ej -1001234567890).'),
 ('persona_1_user_id','<PEGAR>','Tu user_id de Telegram (numerico).'),
 ('persona_1_nombre','Lisandro','Como te nombra el bot.'),
 ('persona_2_user_id','<PEGAR>','El user_id de ella.'),
 ('persona_2_nombre','<PEGAR>','Como la nombra el bot.'),
 ('max_arrastres','3','Ciclos sin comprarse antes de descartar un item solo.'),
 ('categorias','almacén,lácteos,limpieza,higiene,bebidas,frescos,otros','No agregar sin tocar el prompt y el validador.'),
 ('mensaje_fijado_id','','message_id del mensaje de lista fijado en el grupo. Lo escribe el bot.'),
 ('usar_ia','true','false = solo diccionario, cero llamadas al LLM.'),
 ('umbral_confianza','0.4','Por debajo de esto el item se descarta.'),
]
for k, v, d in CONFIG:
    cfg.append([k, v, d])
for r in range(2, len(CONFIG) + 2):
    cfg.cell(row=r, column=1).font = Font(bold=True)
    cfg.cell(row=r, column=3).font = Font(color='6C757D', size=9)
cfg.conditional_formatting.add(f'A2:C{len(CONFIG)+1}', FormulaRule(
    formula=['$B2="<PEGAR>"'], fill=PatternFill('solid', fgColor='FFE0B2'), stopIfTrue=False))

# ---------------- Log ----------------
sheet(wb, 'Log', ['ts','update_id','user_id','nombre','tipo','texto','comando',
                  'items_json','resultado','error'],
      [22,14,14,14,10,44,14,44,14,34])

out = '/home/user/biometric-analyzer/tools/compras-casa.xlsx'
wb.save(out)
print('escrito:', out)
print('hojas:', wb.sheetnames)
print('catalogo sembrado:', len(SEED), 'items')
